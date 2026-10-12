"""Model audit: re-runs every integrity check across data, engines and the Excel workbook.

    py -3.12 audit.py        # writes output/audit.json and output/audit_report.md, exits 1 if anything fails

Checks are identities (things that must hold exactly) and consistency tests, not opinions about assumptions.
"""
from __future__ import annotations
import datetime as dt
import json
import subprocess
import sys
from dataclasses import replace

import numpy as np
import pandas as pd

from model.three_statement import load_history, default_assumptions, forecast, SCENARIOS, BASE_YEAR
from model.dcf import dcf, sensitivity, wacc
from model.setup import load_market, valuation_inputs, bridge_for
from model.montecarlo import MCSpec, make_base, base_path_value
from model.segments import load_segments, SEGMENTS
from model.lbo import run as lbo_run, N_PROJ
from model.lbo_setup import default_inputs as lbo_defaults
from model.ma import run_deal, DealInputs
from model.pipeline import Settings
from model import portfolio as pm

RESULTS: list[dict] = []


def check(area: str, name: str, ok: bool, detail: str = "") -> None:
    RESULTS.append(dict(area=area, check=name, passed=bool(ok), detail=detail))


# ── 1. source data ──────────────────────────────────────────────────────────────
for co in ("shell", "bp"):
    h = load_history(co)
    for ident, lhs, rhs in [("assets = liabilities + equity", h.loc["total_assets"], h.loc["total_liabilities"] + h.loc["equity_total"]),
                            ("equity = shareholders + minorities", h.loc["equity_total"], h.loc["equity_parent"] + h.loc["nci"]),
                            ("pre-tax profit − tax = net income", h.loc["pbt"] - h.loc["tax"], h.loc["net_income"])]:
        err = float((lhs - rhs).abs().max())
        check("Data", f"{co.upper()} FY2021–25: {ident}", err < 0.5, f"max difference ${err:,.1f}m")
    core = ["revenue", "pbt", "tax", "net_income", "cfo", "capex", "total_assets", "total_liabilities", "cash", "equity_parent"]
    missing = int(h.loc[core].isna().sum().sum())
    check("Data", f"{co.upper()}: no missing core lines", missing == 0, f"{missing} missing cells")

seg = load_segments()
group = load_history("shell")
worst = 0.0
for y in (2023, 2024, 2025):
    s = seg[seg.year == y]
    tot = s[s.segment == "Total"].set_index("metric").usd_m
    parts = s[s.segment != "Total"].groupby("metric").usd_m.sum()
    worst = max(worst, float((parts - tot).abs().max()))
    ties = {"revenue_third_party": "revenue", "capex": "capex", "dda": "da", "tax": "tax", "income": "net_income"}
    for sm, gm in ties.items():
        check("Data", f"Shell segments {y}: {sm} ties to group accounts", abs(tot[sm] - group.loc[gm, y]) <= 2,
              f"segment total {tot[sm]:,.0f} vs group {group.loc[gm, y]:,.0f}")
check("Data", "Shell segments: every row sums to its reported total", worst <= 2, f"max rounding difference ${worst:,.0f}m")

px = pm.load_prices()
mk = load_market()
last = px.index.max()
age = (pd.Timestamp(dt.date.today()) - last).days
check("Data", "Price history is recent (≤ 7 days old)", age <= 7, f"last price {last.date()}, {age} days old")
check("Data", "Market inputs match the latest price date", mk["shell"]["price_date"] == str(px["Shell"].dropna().index[-1].date()),
      f"market inputs {mk['shell']['price_date']} vs prices {px['Shell'].dropna().index[-1].date()}")
gaps = {c: int(px[c].loc[px[c].first_valid_index():].isna().sum()) for c in pm.BASKET}
check("Data", "Price gaps are holiday-sized only (< 5% of days)", max(gaps.values()) < 0.05 * len(px), f"max gaps {max(gaps.values())} of {len(px)} days")

# ── 2. three-statement forecasts ────────────────────────────────────────────────
for co in ("shell", "bp"):
    h = load_history(co)
    for sc in SCENARIOS:
        out = forecast(h, default_assumptions(h, sc))
        bc = float(out["balance_sheet"].loc["balance_check"].abs().max())
        cash_open = float(h.loc["cash", BASE_YEAR])
        roll = 0.0
        for y in out["cash_flow"].columns:
            roll = max(roll, abs(cash_open + out["cash_flow"].loc["net_change_in_cash", y] - out["balance_sheet"].loc["cash", y]))
            cash_open = out["balance_sheet"].loc["cash", y]
        check("Forecast", f"{co.upper()} {sc}: balance sheet balances every year", bc < 1e-6, f"max imbalance {bc:.1e}")
        check("Forecast", f"{co.upper()} {sc}: cash rolls forward from the cash flow", roll < 1e-6, f"max difference {roll:.1e}")

# ── 3. DCF and Monte Carlo ──────────────────────────────────────────────────────
for co in ("shell", "bp"):
    h = load_history(co)
    a = default_assumptions(h)
    out = forecast(h, a)
    v, b = valuation_inputs(h, mk, co), bridge_for(h, mk, co, a)
    r = dcf(out, v, b)
    br = r["bridge"]
    check("DCF", f"{co.upper()}: equity bridge adds up", abs(br["equity_value"] - sum(x for k, x in br.items() if k != "equity_value")) < 1e-6)
    check("DCF", f"{co.upper()}: EV = PV of flows + PV of terminal value", abs(r["ev"] - r["pv_flows"] - r["pv_terminal"]) < 1e-6)
    s = sensitivity(out, v, b, [wacc(v)], [v.terminal_growth])
    check("DCF", f"{co.upper()}: sensitivity grid centre = point estimate", abs(s.iloc[0, 0] - r["value_per_share_gbp_pence"]) < 1e-6)
    mc0 = base_path_value(make_base(h, a, v, b), MCSpec())
    check("Monte Carlo", f"{co.upper()}: zero-shock path reproduces the DCF", abs(mc0 - r["value_per_share_gbp_pence"]) < 1e-6,
          f"{mc0:,.4f}p vs {r['value_per_share_gbp_pence']:,.4f}p")

# ── 4. LBO ──────────────────────────────────────────────────────────────────────
for segname in SEGMENTS:
    i = lbo_defaults(segname)
    res = lbo_run(i)
    su, p = res["su"], res["projections"]
    check("LBO", f"{segname}: sources = uses", abs(sum(su["sources"].values()) - sum(su["uses"].values())) < 1e-6)
    check("LBO", f"{segname}: returns bridge reconciles", abs(res["returns"]["bridge_check"]) < 1e-6)
    cash, tlb, worst = su["min_cash"], su["sources"]["term_loan_b"], 0.0
    for t in range(1, N_PROJ + 1):
        c = p[t]
        worst = max(worst, abs(cash + c["cfads"] - c["cash_interest"] - c["mandatory_amort"] + c["revolver_draw"] - c["revolver_repay"] - c["sweep"] - c["cash"]))
        cash = c["cash"]
    check("LBO", f"{segname}: cash rolls forward", worst < 1e-4, f"max difference {worst:.1e}")
    check("LBO", f"{segname}: financeability reported", True, "financeable" if res["financeable"] else "flagged NOT financeable (by design for weak divisions)")

# ── 5. merger ───────────────────────────────────────────────────────────────────
S = Settings()
for label, d in [("default terms", DealInputs()), ("all stock, 10% divested", DealInputs(pct_stock=1.0, divest_pct_ebitda=0.10)),
                 ("all cash, 50% premium", DealInputs(pct_stock=0.0, premium=0.5))]:
    dr = run_deal(S, d)
    check("Merger", f"{label}: opening balance sheet balances", abs(dr["opening_check"]) < 1e-6, f"{dr['opening_check']:.1e}")
    check("Merger", f"{label}: pro forma forecast balances", float(dr["pf"]["balance_sheet"].loc["balance_check"].abs().max()) < 1e-6)
    su = dr["su"]
    check("Merger", f"{label}: sources = uses", abs(sum(su["sources"].values()) - sum(su["uses"].values())) < 1e-6)
    pp = dr["ppa"]
    check("Merger", f"{label}: goodwill = price − fair value of net assets", abs(pp["goodwill_created"] - (pp["equity_purchase_price"] - pp["fair_value_net_assets"])) < 1e-6)

# ── 6. portfolio ────────────────────────────────────────────────────────────────
A = list(pm.BASKET)
rw = pm.returns(pm.gbp_prices(px, A), "W")
e = pm.estimate(rw, A)
W = {k: f(e, 0.25) for k, f in pm.STRATEGIES.items()}
check("Portfolio", "All strategies: weights sum to 100% within caps",
      all(abs(w.sum() - 1) < 1e-8 and w.min() >= -1e-10 and w.max() <= 0.25 + 1e-6 for w in W.values()))
mv = pm.port_stats(W["Min variance"], e)["vol"]
check("Portfolio", "Minimum-variance portfolio has the lowest volatility", all(mv <= pm.port_stats(w, e)["vol"] + 1e-6 for w in W.values()))
ms = pm.port_stats(W["Max Sharpe"], e)["sharpe"]
check("Portfolio", "Maximum-Sharpe portfolio has the highest Sharpe ratio", all(ms >= pm.port_stats(w, e)["sharpe"] - 1e-6 for w in W.values()))
bt1 = pm.backtest(rw, A, cap=0.25, strategies=["Min variance"])
t2 = rw.copy(); t2.iloc[-20:] *= 5
bt2 = pm.backtest(t2, A, cap=0.25, strategies=["Min variance"])
early = [x for x in bt1["weights"]["Min variance"].index if x < rw.index[-20]]
check("Portfolio", "Backtest uses no future data", np.allclose(bt1["weights"]["Min variance"].loc[early].values, bt2["weights"]["Min variance"].loc[early].values))

# ── 7. Excel workbook ───────────────────────────────────────────────────────────
try:
    from openpyxl import load_workbook
    path = "output/XAVI_valuation_model.xlsx"
    wv, wf = load_workbook(path, data_only=True), load_workbook(path)
    n_f = n_err = n_none = 0
    for ws in wf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_f += 1
                    x = wv[ws.title][c.coordinate].value
                    n_err += isinstance(x, str) and x.startswith("#")
                    n_none += x is None
    check("Excel", "Workbook has no formula errors", n_err == 0 and n_none == 0, f"{n_f:,} formulas, {n_err} errors, {n_none} uncalculated")
    ck = wv["Checks"]
    statuses = [ck.cell(r, 6).value for r in range(5, ck.max_row + 1) if ck.cell(r, 6).value in ("PASS", "FAIL", "ERROR")]
    check("Excel", "Workbook reconciles to the Python engine", statuses and all(s_ == "PASS" for s_ in statuses),
          f"{statuses.count('PASS')} of {len(statuses)} reconciliation lines pass")
except FileNotFoundError:
    check("Excel", "Workbook exists", False, "run export_excel.py, then recalc_excel.ps1")

# ── 8. test suite ───────────────────────────────────────────────────────────────
pt = subprocess.run([sys.executable, "-m", "pytest", "-q", "--no-header", "-p", "no:cacheprovider"], capture_output=True, text=True)
summary = [ln for ln in pt.stdout.splitlines() if " passed" in ln or " failed" in ln]
check("Tests", "Automated test suite passes", pt.returncode == 0, summary[-1] if summary else pt.stdout[-200:])

# ── report ──────────────────────────────────────────────────────────────────────
df = pd.DataFrame(RESULTS)
n_ok, n = int(df.passed.sum()), len(df)
stamp = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
json.dump(dict(run_at=stamp, passed=n_ok, total=n, results=RESULTS), open("output/audit.json", "w"), indent=2)
lines = [f"# XAVI model audit", "", f"Run {stamp}: **{n_ok} of {n} checks pass.**", ""]
for area, g in df.groupby("area", sort=False):
    lines += [f"## {area}", "", "| Status | Check | Detail |", "|---|---|---|"]
    lines += [f"| {'PASS' if r_.passed else 'FAIL'} | {r_.check} | {r_.detail} |" for r_ in g.itertuples()]
    lines.append("")
open("output/audit_report.md", "w", encoding="utf8").write("\n".join(lines))
for r_ in df.itertuples():
    if not r_.passed:
        print("FAIL", r_.area, "|", r_.check, "|", r_.detail)
print(f"{n_ok}/{n} checks pass  ->  output/audit_report.md")
sys.exit(0 if n_ok == n else 1)
