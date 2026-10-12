"""Export the XAVI models to an Excel workbook with live formulas.

    py -3.12 export_excel.py            # writes output/XAVI_valuation_model.xlsx
    powershell -File recalc_excel.ps1   # recalculates it in Excel (iterative calc for the interest circularity)
    py -3.12 verify_excel.py            # compares every Excel result with the Python engine

Conventions: blue = hard-coded input, black = formula, green = link to another sheet, yellow fill = key assumption.
The Checks sheet reconciles the workbook to the Python engine and shows PASS/FAIL for each line.
"""
from __future__ import annotations
import datetime as dt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as COL
from openpyxl.workbook.properties import CalcProperties

from model.pipeline import Settings, compute
from model.three_statement import BASE_YEAR, N_YEARS
from model.lbo import run as lbo_run
from model.lbo_setup import default_inputs as lbo_defaults
from model.ma import run_deal, DealInputs
from model import portfolio as pm

OUT = "output/XAVI_valuation_model.xlsx"
AR = "Arial"
F_IN = Font(name=AR, size=10, color="0000FF")
F_LINK = Font(name=AR, size=10, color="008000")
F_CALC = Font(name=AR, size=10, color="000000")
F_BOLD = Font(name=AR, size=10, bold=True)
F_NOTE = Font(name=AR, size=9, italic=True, color="666666")
F_TITLE = Font(name=AR, size=16, bold=True, color="111114")
F_SUB = Font(name=AR, size=11, bold=True, color="8A6D2B")
F_HEAD = Font(name=AR, size=10, bold=True, color="C9A961")
FILL_HEAD = PatternFill("solid", fgColor="111114")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_SECTION = PatternFill("solid", fgColor="F3EEE2")
THIN = Side(style="thin", color="BFBFBF")
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
NUM4 = '#,##0.0000;(#,##0.0000);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
MULT = '0.00"x";(0.00"x");"-"'
YEARS = list(range(BASE_YEAR + 1, BASE_YEAR + 1 + N_YEARS))   # 2026..2035
FC = [COL(3 + i) for i in range(N_YEARS)]                      # C..L forecast columns


def cell(ws, ref, value, font=F_CALC, fmt=None, fill=None, bold=False, comment=None, align=None):
    c = ws[ref]
    c.value = value
    c.font = Font(name=font.name, size=font.size, color=font.color, bold=bold or font.bold, italic=font.italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if comment:
        c.comment = Comment(comment, "XAVI")
    if align:
        c.alignment = Alignment(horizontal=align)
    return c


def header(ws, row, labels: list[str], start_col: int = 1):
    for j, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + j, value=lab)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(horizontal="left" if j == 0 else "right")


def title(ws, text, sub):
    cell(ws, "A1", text, F_TITLE)
    cell(ws, "A2", sub, F_NOTE)


def section(ws, row, text, ncols=12):
    for j in range(1, ncols + 1):
        ws.cell(row=row, column=j).fill = FILL_SECTION
    cell(ws, f"A{row}", text, F_SUB)


def widths(ws, first=46, rest=13, n=14):
    ws.column_dimensions["A"].width = first
    for j in range(2, n + 1):
        ws.column_dimensions[COL(j)].width = rest


# ═════════════════════════════ data from the Python engine ═════════════════════════════
S = Settings()
R = compute(S)                       # Shell, base case
h, a, out, v, b, r = R["h"], R["a"], R["out"], R["v"], R["b"], R["r"]
mk = R["market"]
fx = R["common"]["gbpusd"]
LBO_I = lbo_defaults("Marketing")
LBO_R = lbo_run(LBO_I)
DEAL = run_deal(S, DealInputs())
PX = pm.load_prices()
ASSETS = list(pm.BASKET)
RW = pm.returns(pm.gbp_prices(PX, ASSETS), "W")
EST = pm.estimate(RW, ASSETS, rf=float(R["common"]["short_rate"]), erp=0.05, hist_weight=0.25)
PW = {name: f(EST, 0.25) for name, f in pm.STRATEGIES.items()}

wb = Workbook()
wb.calculation = CalcProperties(iterate=True, iterateCount=10000, iterateDelta=1e-10, fullCalcOnLoad=True)

# ═════════════════════════════ History ═════════════════════════════
wsH = wb.active
wsH.title = "History"
title(wsH, "Shell plc · reported history", "USD millions. Source: Shell plc Form 20-F, FY2021–FY2025, XBRL facts via SEC EDGAR "
      "(https://data.sec.gov/api/xbrl/companyfacts/CIK0001306965.json). Derived lines are noted.")
HY = [2021, 2022, 2023, 2024, 2025]
header(wsH, 4, ["Line item (USD m)"] + [str(y) for y in HY])
HROW = {}
labels = {"revenue": "Revenue", "pbt": "Profit before tax", "tax": "Taxation", "net_income": "Net income",
          "net_income_parent": "Net income to shareholders", "da": "D&A incl. impairments (cash-flow add-back)",
          "finance_costs": "Interest expense", "cfo": "Cash from operations", "cfi": "Cash from investing", "cff": "Cash from financing",
          "capex": "Capital expenditure", "dividends_paid": "Dividends paid", "buybacks": "Share buybacks", "total_assets": "Total assets",
          "current_assets": "Current assets", "cash": "Cash and equivalents", "inventories": "Inventories", "receivables": "Trade and other receivables",
          "ppe": "Property, plant & equipment incl. right-of-use", "goodwill": "Goodwill", "intangibles": "Other intangibles",
          "total_liabilities": "Total liabilities", "current_liabilities": "Current liabilities", "payables": "Trade and other payables",
          "debt_current": "Current borrowings", "debt_long": "Long-term borrowings", "lease_liabilities": "Lease liabilities",
          "equity_total": "Total equity", "equity_parent": "Equity attributable to shareholders", "nci": "Non-controlling interests",
          "net_debt_reported": "Net debt (as reported)", "finance_income": "Interest and other income", "jv_share": "Share of JV/associate profit",
          "equity_investments": "Equity-accounted investments", "impairment_net": "Net impairments (derived)",
          "da_underlying": "D&A excl. impairments (derived)", "lease_principal": "Lease principal repaid (derived)",
          "provisions_debt_like": "Decommissioning provisions", "pension_deficit": "Net pension deficit (derived)",
          "ebit_ops": "Operating EBIT, ex JV and finance items (derived)"}
row = 5
for k, lab in labels.items():
    if k not in h.index:
        continue
    cell(wsH, f"A{row}", lab)
    for j, y in enumerate(HY):
        val = h.loc[k, y]
        cell(wsH, f"{COL(2 + j)}{row}", None if pd.isna(val) else float(val), F_IN, NUM)
    HROW[k] = row
    row += 1
cell(wsH, f"A{row + 1}", "Balance check: total assets − total liabilities − total equity", F_BOLD)
for j in range(5):
    c = COL(2 + j)
    cell(wsH, f"{c}{row + 1}", f"={c}{HROW['total_assets']}-{c}{HROW['total_liabilities']}-{c}{HROW['equity_total']}", F_CALC, NUM)
HCHECK = row + 1
widths(wsH, 52, 13, 7)
wsH.freeze_panes = "B5"
H25 = lambda k: f"History!$F${HROW[k]}"      # FY2025 column

# ═════════════════════════════ Inputs ═════════════════════════════
wsI = wb.create_sheet("Inputs", 0)
title(wsI, "Shell plc · forecast and valuation inputs", "Base case. Blue cells are inputs; change them and every sheet recalculates. "
      "Defaults are calibrated from Shell's own history by the Python engine (model/three_statement.py).")
header(wsI, 4, ["Annual drivers", "Unit"] + [str(y) for y in YEARS])
IROW = {}
vec = [("revenue_growth", "Revenue growth", "%", a.revenue_growth, PCT, "Placeholder macro view: -1%, +1%, +2%... fading to terminal growth."),
       ("ebit_margin", "Operating EBIT margin", "%", a.ebit_margin, PCT, "2023-25 average margin excluding impairments, less a 1% recurring impairment allowance."),
       ("da_pct", "D&A excl. impairments", "% rev", a.da_pct, PCT2, "FY2025 ratio."),
       ("capex_pct", "Capital expenditure", "% rev", a.capex_pct, PCT2, "FY2025 ratio."),
       ("lease_pct", "New leases = lease principal repaid", "% rev", a.lease_pct, PCT2, "FY2025 ratio; keeps the lease liability flat."),
       ("tax_rate", "Tax rate on pre-tax profit ex JV", "%", a.tax_rate, PCT, "2023-25 effective rate ex impairments, fading 2.5pp over 10 years."),
       ("jv_pct", "Share of JV profit", "% rev", a.jv_pct, PCT2, "2023-25 average; assumed fully distributed in cash."),
       ("dividends", "Dividends", "USD m", a.dividends, NUM, "FY2025 dividends growing 4% a year to 2030, 2.5% after."),
       ("buybacks", "Buyback target (flexes to protect cash floor)", "USD m", a.buybacks, NUM, "FY2025 level."),
       ("net_debt_issuance", "Net new borrowing", "USD m", a.net_debt_issuance, NUM, "None assumed.")]
row = 5
for key, lab, unit, vals, fmt, note in vec:
    cell(wsI, f"A{row}", lab, comment=note)
    cell(wsI, f"B{row}", unit, F_NOTE)
    for j, x in enumerate(vals):
        cell(wsI, f"{FC[j]}{row}", float(x), F_IN, fmt, fill=FILL_KEY if key in ("revenue_growth", "ebit_margin") else None)
    IROW[key] = row
    row += 1
row += 1
section(wsI, row, "Balance-sheet and financing drivers")
row += 1
scal = [("receivable_pct", "Receivables", "% rev", a.receivable_pct, PCT), ("inventory_pct", "Inventories", "% rev", a.inventory_pct, PCT),
        ("payable_pct", "Payables", "% rev", a.payable_pct, PCT), ("finance_rate", "Interest rate on borrowings + leases", "%", a.finance_rate, PCT2),
        ("cash_rate", "Interest rate on cash", "%", a.cash_rate, PCT2), ("nci_income_fixed", "Minority (NCI) income, fixed", "USD m", a.nci_income_fixed, NUM),
        ("min_cash_pct", "Minimum cash", "% rev", a.min_cash_pct, PCT),
        ("circ", "Circularity switch: 1 = interest on average balances, 0 = opening balances", "1/0", 1, "0")]
for key, lab, unit, x, fmt in scal:
    cell(wsI, f"A{row}", lab); cell(wsI, f"B{row}", unit, F_NOTE); cell(wsI, f"C{row}", float(x), F_IN, fmt)
    IROW[key] = row
    row += 1
row += 1
section(wsI, row, "Valuation inputs")
row += 1
val_inputs = [("risk_free", "Risk-free rate (US 10y, blend of spot and 1-year average)", "%", v.risk_free, PCT2, "Yahoo Finance ^TNX; spot and 1y average blended."),
              ("erp", "Equity risk premium", "%", v.equity_risk_premium, PCT2, "Assumption."),
              ("beta", "Beta (5y weekly vs FTSE 100, Blume-adjusted)", "x", v.beta, NUM2, "Computed from Yahoo Finance prices."),
              ("credit_spread", "Credit spread over risk-free", "%", v.credit_spread, PCT2, "Assumption."),
              ("tax_shield", "Tax rate on interest / provisions", "%", v.tax_shield_rate, PCT, "Assumption (marginal rate)."),
              ("terminal_growth", "Terminal growth", "%", v.terminal_growth, PCT2, "Assumption."),
              ("exit_multiple", "Exit EV/EBITDA (cross-check)", "x", v.exit_multiple, MULT, "Assumption."),
              ("shares_m", "Shares outstanding", "m", mk["shares_m"], NUM1, "Yahoo Finance; matches the 20-F cover (5,690m at 31 Dec 2025)."),
              ("price_pence", "Share price", "GBp", mk["price_gbp_pence"], NUM1, f"London close {mk['price_date']}, Yahoo Finance."),
              ("gbpusd", "GBP/USD", "USD per GBP", fx, NUM4, "Yahoo Finance."),
              ("valuation_date", "Valuation date", "date", dt.datetime(2026, 10, 10), "yyyy-mm-dd", "Value struck at 31 Dec 2025, rolled forward to this date.")]
for key, lab, unit, x, fmt, note in val_inputs:
    cell(wsI, f"A{row}", lab, comment=note); cell(wsI, f"B{row}", unit, F_NOTE)
    cell(wsI, f"C{row}", x if isinstance(x, dt.datetime) else float(x), F_IN, fmt, fill=FILL_KEY if key in ("erp", "terminal_growth") else None)
    IROW[key] = row
    row += 1
IN = lambda k, c="C": f"Inputs!${c}${IROW[k]}"
INV = lambda k, c: f"Inputs!{c}${IROW[k]}"
widths(wsI, 56, 11, 13)
wsI.freeze_panes = "C5"

# ═════════════════════════════ Forecast ═════════════════════════════
wsF = wb.create_sheet("Forecast", 1)
title(wsF, "Shell plc · integrated three-statement forecast", "USD millions. Column B is the FY2025 opening balance sheet (from History); "
      "C–L are forecasts. Interest is on average balances, so the sheet is circular: iterative calculation is switched on.")
header(wsF, 4, ["", "2025A"] + [f"{y}E" for y in YEARS])
FR: dict[str, int] = {}
lines = ["IS", "revenue", "ebit", "da", "ebitda", "interest_expense", "interest_income", "pbt_ex_jv", "tax", "jv_income", "net_income",
         "nci_income", "net_income_parent", "",
         "WC", "receivables", "inventories", "payables", "d_nwc", "min_cash", "",
         "CF", "cfo", "capex", "lease_add", "cfi", "pre_cash", "avail", "draw", "repay", "buybacks", "dividends", "cff", "net_change", "",
         "BS", "cash", "fixed_assets", "goodwill", "other_assets", "total_assets", "borrowings", "revolver", "leases", "other_liabilities",
         "total_liabilities", "equity_parent", "nci", "total_equity", "balance_check", "",
         "KP", "ufcf", "net_debt", "nd_ebitda"]
names = {"IS": "Income statement", "WC": "Working capital", "CF": "Cash flow", "BS": "Balance sheet", "KP": "Free cash flow and credit",
         "revenue": "Revenue", "ebit": "Operating EBIT", "da": "D&A", "ebitda": "EBITDA", "interest_expense": "Interest expense",
         "interest_income": "Interest income", "pbt_ex_jv": "Pre-tax profit ex JV", "tax": "Tax", "jv_income": "Share of JV profit",
         "net_income": "Net income", "nci_income": "Minorities", "net_income_parent": "Net income to shareholders",
         "receivables": "Receivables", "inventories": "Inventories", "payables": "Payables", "d_nwc": "Increase in working capital",
         "min_cash": "Minimum cash (floor)", "cfo": "Cash from operations", "capex": "Capital expenditure", "lease_add": "Lease principal / new leases",
         "cfi": "Cash from investing", "pre_cash": "Cash before buybacks and revolver", "avail": "Cash available above the floor",
         "draw": "Revolver drawn", "repay": "Revolver repaid", "buybacks": "Share buybacks", "dividends": "Dividends", "cff": "Cash from financing",
         "net_change": "Net change in cash", "cash": "Cash", "fixed_assets": "Fixed assets incl. right-of-use", "goodwill": "Goodwill",
         "other_assets": "Other assets (plug, held flat)", "total_assets": "Total assets", "borrowings": "Borrowings", "revolver": "Revolver",
         "leases": "Lease liabilities", "other_liabilities": "Other liabilities (plug, held flat)", "total_liabilities": "Total liabilities",
         "equity_parent": "Shareholders’ equity", "nci": "Minorities", "total_equity": "Total equity", "balance_check": "Balance check (must be 0)",
         "ufcf": "Unlevered free cash flow", "net_debt": "Net debt incl. leases", "nd_ebitda": "Net debt / EBITDA"}
row = 5
for k in lines:
    if k == "":
        row += 1
        continue
    if k in ("IS", "WC", "CF", "BS", "KP"):
        section(wsF, row, names[k])
    else:
        cell(wsF, f"A{row}", names[k], bold=k in ("revenue", "ebitda", "net_income", "total_assets", "total_liabilities", "total_equity", "ufcf"))
        FR[k] = row
    row += 1
X = lambda k, c: f"{c}{FR[k]}"

# opening balances (column B), linked to History
opening = {
    "revenue": f"={H25('revenue')}", "cash": f"={H25('cash')}", "receivables": f"={H25('receivables')}", "inventories": f"={H25('inventories')}",
    "payables": f"={H25('payables')}", "fixed_assets": f"={H25('ppe')}+{H25('intangibles')}", "goodwill": f"={H25('goodwill')}",
    "borrowings": f"={H25('debt_current')}+{H25('debt_long')}", "revolver": 0, "leases": f"={H25('lease_liabilities')}",
    "equity_parent": f"={H25('equity_parent')}", "nci": f"={H25('nci')}",
}
for k, f in opening.items():
    cell(wsF, X(k, "B"), f, F_LINK if isinstance(f, str) else F_IN, NUM)
cell(wsF, X("other_assets", "B"), f"={H25('total_assets')}-B{FR['cash']}-B{FR['receivables']}-B{FR['inventories']}-B{FR['fixed_assets']}-B{FR['goodwill']}",
     F_LINK, NUM, comment="Plug: lines the forecast does not drive (JV stakes, derivatives, deferred tax...). Held flat.")
cell(wsF, X("other_liabilities", "B"), f"={H25('total_liabilities')}-B{FR['payables']}-B{FR['borrowings']}-B{FR['leases']}",
     F_LINK, NUM, comment="Plug: provisions, deferred tax, derivatives... Held flat.")
cell(wsF, X("total_assets", "B"), f"=B{FR['cash']}+B{FR['receivables']}+B{FR['inventories']}+B{FR['fixed_assets']}+B{FR['goodwill']}+B{FR['other_assets']}", F_CALC, NUM)
cell(wsF, X("total_liabilities", "B"), f"=B{FR['payables']}+B{FR['borrowings']}+B{FR['revolver']}+B{FR['leases']}+B{FR['other_liabilities']}", F_CALC, NUM)
cell(wsF, X("total_equity", "B"), f"=B{FR['equity_parent']}+B{FR['nci']}", F_CALC, NUM)
cell(wsF, X("balance_check", "B"), f"=B{FR['total_assets']}-B{FR['total_liabilities']}-B{FR['total_equity']}", F_CALC, NUM)

prevc = ["B"] + FC[:-1]
for c, p in zip(FC, prevc):
    g = {}
    g["revenue"] = f"={p}{FR['revenue']}*(1+{INV('revenue_growth', c)})"
    g["ebit"] = f"={c}{FR['revenue']}*{INV('ebit_margin', c)}"
    g["da"] = f"={c}{FR['revenue']}*{INV('da_pct', c)}"
    g["ebitda"] = f"={c}{FR['ebit']}+{c}{FR['da']}"
    debt_end = f"({p}{FR['borrowings']}+{p}{FR['revolver']}+{INV('net_debt_issuance', c)}+{c}{FR['draw']}-{c}{FR['repay']})"
    open_debt = f"({p}{FR['borrowings']}+{p}{FR['revolver']})"
    g["interest_expense"] = (f"=IFERROR({IN('finance_rate')}*(IF({IN('circ')}=1,({open_debt}+{debt_end})/2,{open_debt})+{p}{FR['leases']}),0)")
    g["interest_income"] = f"=IFERROR({IN('cash_rate')}*IF({IN('circ')}=1,({p}{FR['cash']}+{c}{FR['cash']})/2,{p}{FR['cash']}),0)"
    g["pbt_ex_jv"] = f"={c}{FR['ebit']}-{c}{FR['interest_expense']}+{c}{FR['interest_income']}"
    g["tax"] = f"={c}{FR['pbt_ex_jv']}*{INV('tax_rate', c)}"
    g["jv_income"] = f"={c}{FR['revenue']}*{INV('jv_pct', c)}"
    g["net_income"] = f"={c}{FR['pbt_ex_jv']}-{c}{FR['tax']}+{c}{FR['jv_income']}"
    g["nci_income"] = f"={IN('nci_income_fixed')}"
    g["net_income_parent"] = f"={c}{FR['net_income']}-{c}{FR['nci_income']}"
    g["receivables"] = f"={c}{FR['revenue']}*{IN('receivable_pct')}"
    g["inventories"] = f"={c}{FR['revenue']}*{IN('inventory_pct')}"
    g["payables"] = f"={c}{FR['revenue']}*{IN('payable_pct')}"
    g["d_nwc"] = f"=({c}{FR['receivables']}-{p}{FR['receivables']})+({c}{FR['inventories']}-{p}{FR['inventories']})-({c}{FR['payables']}-{p}{FR['payables']})"
    g["min_cash"] = f"={c}{FR['revenue']}*{IN('min_cash_pct')}"
    g["cfo"] = f"={c}{FR['net_income']}+{c}{FR['da']}-{c}{FR['d_nwc']}"
    g["capex"] = f"={c}{FR['revenue']}*{INV('capex_pct', c)}"
    g["lease_add"] = f"={c}{FR['revenue']}*{INV('lease_pct', c)}"
    g["cfi"] = f"=-{c}{FR['capex']}"
    g["dividends"] = f"={INV('dividends', c)}"
    g["pre_cash"] = f"={p}{FR['cash']}+{c}{FR['cfo']}+{c}{FR['cfi']}+{INV('net_debt_issuance', c)}-{c}{FR['dividends']}-{c}{FR['nci_income']}-{c}{FR['lease_add']}"
    g["avail"] = f"={c}{FR['pre_cash']}-{c}{FR['min_cash']}"
    g["draw"] = f"=MAX(0,-{c}{FR['avail']})"
    g["repay"] = f"=MIN({p}{FR['revolver']},MAX(0,{c}{FR['avail']}))"
    g["buybacks"] = f"=MAX(0,MIN({INV('buybacks', c)},{c}{FR['avail']}-{c}{FR['repay']}))"
    g["cff"] = f"={INV('net_debt_issuance', c)}+{c}{FR['draw']}-{c}{FR['repay']}-{c}{FR['dividends']}-{c}{FR['buybacks']}-{c}{FR['nci_income']}-{c}{FR['lease_add']}"
    g["net_change"] = f"={c}{FR['cfo']}+{c}{FR['cfi']}+{c}{FR['cff']}"
    g["cash"] = f"={p}{FR['cash']}+{c}{FR['net_change']}"
    g["fixed_assets"] = f"={p}{FR['fixed_assets']}+{c}{FR['capex']}+{c}{FR['lease_add']}-{c}{FR['da']}"
    g["goodwill"] = f"={p}{FR['goodwill']}"
    g["other_assets"] = f"={p}{FR['other_assets']}"
    g["total_assets"] = f"={c}{FR['cash']}+{c}{FR['receivables']}+{c}{FR['inventories']}+{c}{FR['fixed_assets']}+{c}{FR['goodwill']}+{c}{FR['other_assets']}"
    g["borrowings"] = f"={p}{FR['borrowings']}+{INV('net_debt_issuance', c)}"
    g["revolver"] = f"={p}{FR['revolver']}+{c}{FR['draw']}-{c}{FR['repay']}"
    g["leases"] = f"={p}{FR['leases']}"
    g["other_liabilities"] = f"={p}{FR['other_liabilities']}"
    g["total_liabilities"] = f"={c}{FR['payables']}+{c}{FR['borrowings']}+{c}{FR['revolver']}+{c}{FR['leases']}+{c}{FR['other_liabilities']}"
    g["equity_parent"] = f"={p}{FR['equity_parent']}+{c}{FR['net_income_parent']}-{c}{FR['dividends']}-{c}{FR['buybacks']}"
    g["nci"] = f"={p}{FR['nci']}"
    g["total_equity"] = f"={c}{FR['equity_parent']}+{c}{FR['nci']}"
    g["balance_check"] = f"={c}{FR['total_assets']}-{c}{FR['total_liabilities']}-{c}{FR['total_equity']}"
    g["ufcf"] = f"={c}{FR['ebit']}*(1-{INV('tax_rate', c)})+{c}{FR['da']}-{c}{FR['capex']}-{c}{FR['lease_add']}-{c}{FR['d_nwc']}"
    g["net_debt"] = f"={c}{FR['borrowings']}+{c}{FR['revolver']}+{c}{FR['leases']}-{c}{FR['cash']}"
    g["nd_ebitda"] = f"=IF({c}{FR['ebitda']}<>0,{c}{FR['net_debt']}/{c}{FR['ebitda']},0)"
    for k, f in g.items():
        cell(wsF, X(k, c), f, F_CALC, MULT if k == "nd_ebitda" else NUM)
widths(wsF, 40, 11, 13)
wsF.freeze_panes = "C5"

# ═════════════════════════════ DCF ═════════════════════════════
wsD = wb.create_sheet("DCF", 2)
title(wsD, "Shell plc · discounted cash flow", "USD millions unless stated. Mid-year discounting from 31 Dec 2025; value rolled forward to the valuation date.")
D = {}
rowsD = [("wacc_h", None), ("ke", "Cost of equity = risk-free + beta × ERP"), ("kd", "After-tax cost of debt = (risk-free + spread) × (1 − tax)"),
         ("mcap", "Market capitalisation (USD m) = shares × price × GBP/USD"), ("debt_book", "Debt incl. leases at book (USD m)"),
         ("wacc", "WACC"), ("blank", None), ("flows_h", None)]
r0 = 4
section(wsD, r0, "Cost of capital"); r_ = r0 + 1
for k, lab, f, fmt in [("ke", "Cost of equity = risk-free + beta × ERP", f"={IN('risk_free')}+{IN('beta')}*{IN('erp')}", PCT2),
                       ("kd", "After-tax cost of debt = (risk-free + spread) × (1 − tax)", f"=({IN('risk_free')}+{IN('credit_spread')})*(1-{IN('tax_shield')})", PCT2),
                       ("mcap", "Market capitalisation (USD m)", f"={IN('shares_m')}*{IN('price_pence')}/100*{IN('gbpusd')}", NUM),
                       ("debt_book", "Borrowings + leases at book (USD m)", f"={H25('debt_current')}+{H25('debt_long')}+{H25('lease_liabilities')}", NUM),
                       ("wacc", "WACC", None, PCT2)]:
    cell(wsD, f"A{r_}", lab, bold=k == "wacc")
    D[k] = r_
    if f:
        cell(wsD, f"B{r_}", f, F_LINK if "History!" in f else F_CALC, fmt)
    r_ += 1
cell(wsD, f"B{D['wacc']}", f"=(B{D['mcap']}*B{D['ke']}+B{D['debt_book']}*B{D['kd']})/(B{D['mcap']}+B{D['debt_book']})", F_CALC, PCT2, bold=True)
r_ += 1
section(wsD, r_, "Cash flows"); r_ += 1
header(wsD, r_, ["", ""] + [f"{y}E" for y in YEARS]); D["yr"] = r_; r_ += 1
for k, lab in [("ufcf", "Unlevered free cash flow"), ("t", "Discount period (years, mid-year)"), ("df", "Discount factor"), ("pv", "Present value")]:
    cell(wsD, f"A{r_}", lab); D[k] = r_; r_ += 1
for j, c in enumerate(FC):
    cell(wsD, f"{c}{D['ufcf']}", f"=Forecast!{c}{FR['ufcf']}", F_LINK, NUM)
    cell(wsD, f"{c}{D['t']}", j + 0.5, F_CALC, NUM1)
    cell(wsD, f"{c}{D['df']}", f"=1/(1+$B${D['wacc']})^{c}{D['t']}", F_CALC, NUM4)
    cell(wsD, f"{c}{D['pv']}", f"={c}{D['ufcf']}*{c}{D['df']}", F_CALC, NUM)
last = FC[-1]
r_ += 1
section(wsD, r_, "Enterprise value to equity value"); r_ += 1
items = [("pv_flows", "PV of 10-year cash flows", f"=SUM(C{D['pv']}:{last}{D['pv']})", NUM),
         ("tv", "Terminal value (Gordon) at 2035", f"={last}{D['ufcf']}*(1+{IN('terminal_growth')})/($B${D['wacc']}-{IN('terminal_growth')})", NUM),
         ("pv_tv", "PV of terminal value", f"=B{{tv}}/(1+$B${D['wacc']})^({last}{D['t']})", NUM),
         ("ev", "Enterprise value", "=B{pv_flows}+B{pv_tv}", NUM),
         ("jv", "+ Equity-accounted JV stakes (book)", f"={H25('equity_investments')}", NUM),
         ("borr", "− Borrowings", f"=-({H25('debt_current')}+{H25('debt_long')})", NUM),
         ("lease", "− Lease liabilities", f"=-{H25('lease_liabilities')}", NUM),
         ("cash", "+ Cash", f"={H25('cash')}", NUM),
         ("nci", "− Minorities", f"=-{H25('nci')}", NUM),
         ("prov", "− Decommissioning provisions, after tax", f"=-{H25('provisions_debt_like')}*(1-{IN('tax_shield')})", NUM),
         ("pens", "− Pension deficit, after tax", f"=-{H25('pension_deficit')}*(1-{IN('tax_shield')})", NUM),
         ("equity", "Equity value at 31 Dec 2025", "=SUM(B{ev}:B{pens})", NUM),
         ("f", "Roll-forward (years from 31 Dec 2025)", f"=({IN('valuation_date')}-DATE(2025,12,31))/365", NUM4),
         ("dps", "Dividend per share paid in the interim (USD)", f"={INV('dividends', 'C')}/{IN('shares_m')}", NUM4),
         ("ps_usd", "Value per share (USD)", "=B{equity}/" + IN('shares_m') + f"*(1+B{D['ke']})^B{{f}}-B{{dps}}*B{{f}}", NUM2),
         ("ps_gbp", "Value per share (pence)", "=B{ps_usd}/" + IN('gbpusd') + "*100", NUM1),
         ("upside", "Upside to market price", "=B{ps_gbp}/" + IN('price_pence') + "-1", PCT),
         ("tv_share", "Terminal value share of EV", "=B{pv_tv}/B{ev}", PCT),
         ("impl_mult", "Implied exit EV/EBITDA", f"=B{{tv}}/Forecast!{last}{FR['ebitda']}", MULT)]
for k, lab, f, fmt in items:
    D[k] = r_; r_ += 1
for k, lab, f, fmt in items:
    ff = f.format(**{kk: D[kk] for kk in D})
    cell(wsD, f"A{D[k]}", lab, bold=k in ("ev", "equity", "ps_gbp"))
    cell(wsD, f"B{D[k]}", ff, F_LINK if ("History!" in ff or "Forecast!" in ff) else F_CALC, fmt, bold=k in ("ev", "equity", "ps_gbp"),
         fill=FILL_KEY if k == "ps_gbp" else None)
r_ += 1
section(wsD, r_, "Cross-check: exit multiple"); r_ += 1
for k, lab, f, fmt in [("tv_m", "Terminal value = exit multiple × 2035 EBITDA", f"={IN('exit_multiple')}*Forecast!{last}{FR['ebitda']}", NUM),
                       ("pv_tv_m", "PV of terminal value (end-year)", f"=B{{tv_m}}/(1+$B${D['wacc']})^10", NUM),
                       ("eq_m", "Equity value", f"=B{D['pv_flows']}+B{{pv_tv_m}}+SUM(B{D['jv']}:B{D['pens']})", NUM),
                       ("ps_m", "Value per share (pence)", f"=(B{{eq_m}}/{IN('shares_m')}*(1+B{D['ke']})^B{D['f']}-B{D['dps']}*B{D['f']})/{IN('gbpusd')}*100", NUM1)]:
    D[k] = r_; r_ += 1
    cell(wsD, f"A{D[k]}", lab)
for k, lab, f, fmt in [("tv_m", "", f"={IN('exit_multiple')}*Forecast!{last}{FR['ebitda']}", NUM),
                       ("pv_tv_m", "", f"=B{{tv_m}}/(1+$B${D['wacc']})^10", NUM),
                       ("eq_m", "", f"=B{D['pv_flows']}+B{{pv_tv_m}}+SUM(B{D['jv']}:B{D['pens']})", NUM),
                       ("ps_m", "", f"=(B{{eq_m}}/{IN('shares_m')}*(1+B{D['ke']})^B{D['f']}-B{D['dps']}*B{D['f']})/{IN('gbpusd')}*100", NUM1)]:
    ff = f.format(**{kk: D[kk] for kk in D})
    cell(wsD, f"B{D[k]}", ff, F_LINK if "Forecast!" in ff else F_CALC, fmt)
r_ += 1
section(wsD, r_, "Sensitivity: value per share (pence), WACC down the side, terminal growth across"); r_ += 1
D["sens_head"] = r_
g_off = [-0.01, -0.005, 0.0, 0.005, 0.01]
cell(wsD, f"A{r_}", "WACC \\ g", F_BOLD)
for j, dg in enumerate(g_off):
    cell(wsD, f"{COL(2 + j)}{r_}", f"={IN('terminal_growth')}+{dg}", F_CALC, PCT, bold=True)
for i, dw in enumerate(g_off):
    rr = r_ + 1 + i
    cell(wsD, f"A{rr}", f"=$B${D['wacc']}+{dw}", F_CALC, PCT2, bold=True)
    for j in range(5):
        gc = f"{COL(2 + j)}${r_}"
        wc = f"$A{rr}"
        pvs = f"SUMPRODUCT($C${D['ufcf']}:${last}${D['ufcf']}/(1+{wc})^$C${D['t']}:${last}${D['t']})"
        tvf = f"${last}${D['ufcf']}*(1+{gc})/({wc}-{gc})/(1+{wc})^${last}${D['t']}"
        f = (f"=(({pvs}+{tvf}+SUM($B${D['jv']}:$B${D['pens']}))/{IN('shares_m')}*(1+$B${D['ke']})^$B${D['f']}-$B${D['dps']}*$B${D['f']})"
             f"/{IN('gbpusd')}*100")
        cell(wsD, f"{COL(2 + j)}{rr}", f, F_CALC, NUM)
D["sens_centre"] = f"D{r_ + 3}"
widths(wsD, 52, 12, 13)

# ═════════════════════════════ LBO ═════════════════════════════
wsL = wb.create_sheet("LBO", 3)
title(wsL, "Carve-out LBO · Shell Marketing division", "USD millions. Division financials from Shell's FY2025 Form 20-F, note 7 (segment information). "
      "Sweep, revolver and interest on average balances are circular: iterative calculation is on.")
i = LBO_I
LI = {}
sec_rows = [("Operating base", [("revenue0", "LTM revenue incl. inter-segment", i.revenue0, NUM), ("ltm_ebitda", "LTM EBITDA (reported basis)", i.ltm_ebitda, NUM),
                                ("da_pct", "D&A excl. impairments, % revenue", i.da_pct, PCT2), ("capex_pct", "Capex, % revenue", i.capex_pct, PCT2),
                                ("nwc_pct", "Working capital, % revenue (group proxy)", i.nwc_pct, PCT2), ("sa_pct", "Standalone costs, % revenue", i.standalone_cost_pct, PCT2),
                                ("tax", "Tax rate", i.tax_rate, PCT), ("cap", "Interest deductibility cap, % EBITDA", i.interest_cap_pct_ebitda, PCT)]),
            ("Deal", [("entry", "Entry EV / LTM EBITDA", i.entry_multiple, MULT), ("exit", "Exit EV / EBITDA", i.exit_mult, MULT),
                      ("hold", "Hold period (years)", i.hold_years, "0"), ("adv", "Advisory fees, % EV", i.advisory_fee_pct, PCT2),
                      ("sep", "Separation costs, % EV", i.separation_cost_pct, PCT2), ("exit_cost", "Exit costs, % exit EV", i.exit_cost_pct, PCT2),
                      ("mincash", "Minimum cash, % revenue", i.min_cash_pct_revenue, PCT2), ("promote", "Management promote, % of gain", i.mgmt_promote_pct, PCT)]),
            ("Financing", [("tlb_x", "Term loan B, x pro-forma EBITDA", i.tlb_x, MULT), ("notes_x", "Senior notes, x", i.notes_x, MULT), ("mezz_x", "Mezzanine, x", i.mezz_x, MULT),
                           ("ltv", "Maximum debt, % of EV", i.max_ltv, PCT), ("base", "Base rate (US 13-week T-bill)", i.base_rate, PCT2),
                           ("tlb_sp", "TLB spread", i.tlb_spread, PCT2), ("coupon", "Notes coupon", i.notes_coupon, PCT2), ("mz_cash", "Mezzanine cash coupon", i.mezz_cash, PCT2),
                           ("mz_pik", "Mezzanine PIK coupon", i.mezz_pik, PCT2), ("rcf_x", "Revolver commitment, x", i.revolver_commit_x, MULT),
                           ("rcf_sp", "Revolver spread", i.revolver_spread, PCT2), ("undrawn", "Undrawn fee", i.undrawn_fee, PCT2), ("amort", "TLB amortisation, % original", i.tlb_amort_pct, PCT2),
                           ("sweep", "Cash sweep", i.sweep_pct, PCT), ("finfee", "Financing fees, % funded debt", i.financing_fee_pct, PCT2), ("feeyrs", "Fee amortisation (years)", i.fee_amort_years, "0"),
                           ("circ", "Circularity switch: 1 = average balances, 0 = opening", 1, "0")])]
r_ = 4
for sec_name, rows_ in sec_rows:
    section(wsL, r_, sec_name); r_ += 1
    for k, lab, x, fmt in rows_:
        cell(wsL, f"A{r_}", lab); cell(wsL, f"B{r_}", float(x), F_IN, fmt, fill=FILL_KEY if k in ("entry", "exit", "tlb_x") else None)
        LI[k] = r_; r_ += 1
P = lambda k: f"$B${LI[k]}"
r_ += 1
section(wsL, r_, "Sources and uses"); r_ += 1
su_rows = [("ev", "Purchase price (EV)", f"={P('entry')}*{P('ltm_ebitda')}"), ("pf", "Pro-forma EBITDA (after standalone costs)", f"={P('ltm_ebitda')}-{P('revenue0')}*{P('sa_pct')}"),
           ("scale", "Debt scale factor (LTV cap)", None), ("tlb0", "Term loan B", None), ("notes", "Senior notes", None), ("mezz0", "Mezzanine", None),
           ("debt", "Total debt", None), ("rcf", "Revolver commitment (undrawn)", f"={P('rcf_x')}*B{{pf}}"), ("finfees", "Financing fees", None),
           ("advfees", "Advisory fees", f"={P('adv')}*B{{ev}}"), ("sepcost", "Separation costs", f"={P('sep')}*B{{ev}}"), ("cash0", "Cash to balance sheet", f"={P('mincash')}*{P('revenue0')}"),
           ("uses", "Total uses", None), ("equity", "Sponsor equity", None)]
for k, lab, f in su_rows:
    LI[k] = r_; r_ += 1
fx_su = {"scale": f"=MIN(1,{P('ltv')}*B{LI['ev']}/(({P('tlb_x')}+{P('notes_x')}+{P('mezz_x')})*B{LI['pf']}))",
         "tlb0": f"={P('tlb_x')}*B{LI['pf']}*B{LI['scale']}", "notes": f"={P('notes_x')}*B{LI['pf']}*B{LI['scale']}",
         "mezz0": f"={P('mezz_x')}*B{LI['pf']}*B{LI['scale']}", "debt": f"=B{LI['tlb0']}+B{LI['notes']}+B{LI['mezz0']}",
         "finfees": f"={P('finfee')}*B{LI['debt']}+0.01*B{LI['rcf']}", "uses": f"=B{LI['ev']}+B{LI['advfees']}+B{LI['finfees']}+B{LI['sepcost']}+B{LI['cash0']}",
         "equity": f"=B{LI['uses']}-B{LI['debt']}"}
for k, lab, f in su_rows:
    ff = fx_su.get(k, f).format(**{kk: LI[kk] for kk in LI})
    cell(wsL, f"A{LI[k]}", lab, bold=k in ("ev", "debt", "uses", "equity"))
    cell(wsL, f"B{LI[k]}", ff, F_CALC, NUM4 if k == "scale" else NUM, bold=k in ("ev", "debt", "uses", "equity"))
r_ += 1
section(wsL, r_, "Annual projection"); r_ += 1
LY = [COL(3 + t) for t in range(7)]       # C..I = years 1..7, B = close
header(wsL, r_, ["", "Close"] + [f"Year {t}" for t in range(1, 8)]); r_ += 1
lrows = ["growth", "margin", "revenue", "ebitda", "da", "ebit", "capex", "dnwc", "min_cash", "mand", "int_rcf", "int_tlb", "int_notes", "int_mezz",
         "pik", "int_inc", "cash_int", "gross_int", "deduct", "tax", "cfads", "avail", "draw", "repay", "sweep", "fcf",
         "cash", "revolver", "tlb", "notes_b", "mezz", "total_debt", "net_debt", "lev", "cover"]
lnames = {"growth": "Revenue growth (input)", "margin": "EBITDA margin before standalone costs (input)", "revenue": "Revenue", "ebitda": "EBITDA (after standalone costs)",
          "da": "D&A", "ebit": "EBIT", "capex": "Capex", "dnwc": "Increase in working capital", "min_cash": "Minimum cash", "mand": "Mandatory TLB amortisation",
          "int_rcf": "Revolver interest + undrawn fee", "int_tlb": "TLB interest", "int_notes": "Notes interest", "int_mezz": "Mezzanine cash interest",
          "pik": "Mezzanine PIK accrual", "int_inc": "Interest income on cash", "cash_int": "Net cash interest", "gross_int": "Net interest incl. PIK and fee amortisation",
          "deduct": "Deductible interest (capped)", "tax": "Cash tax", "cfads": "Cash flow before debt service", "avail": "Cash available above the floor",
          "draw": "Revolver drawn", "repay": "Revolver repaid", "sweep": "Cash sweep to TLB", "fcf": "Free cash flow after interest", "cash": "Cash",
          "revolver": "Revolver", "tlb": "Term loan B", "notes_b": "Senior notes", "mezz": "Mezzanine (with PIK)", "total_debt": "Total debt",
          "net_debt": "Net debt", "lev": "Debt / EBITDA", "cover": "EBITDA / net cash interest"}
LR = {}
for k in lrows:
    cell(wsL, f"A{r_}", lnames[k], bold=k in ("ebitda", "cfads", "total_debt"))
    LR[k] = r_; r_ += 1
for t, c in enumerate(LY):
    cell(wsL, f"{c}{LR['growth']}", float(i.revenue_growth[t]), F_IN, PCT)
    cell(wsL, f"{c}{LR['margin']}", float(i.ebitda_margin[t]), F_IN, PCT2)
cell(wsL, f"B{LR['revenue']}", f"={P('revenue0')}", F_CALC, NUM)
cell(wsL, f"B{LR['cash']}", f"=B{LI['cash0']}", F_CALC, NUM)
cell(wsL, f"B{LR['revolver']}", 0, F_IN, NUM)
cell(wsL, f"B{LR['tlb']}", f"=B{LI['tlb0']}", F_CALC, NUM)
cell(wsL, f"B{LR['notes_b']}", f"=B{LI['notes']}", F_CALC, NUM)
cell(wsL, f"B{LR['mezz']}", f"=B{LI['mezz0']}", F_CALC, NUM)
cell(wsL, f"B{LR['total_debt']}", f"=B{LR['revolver']}+B{LR['tlb']}+B{LR['notes_b']}+B{LR['mezz']}", F_CALC, NUM)
cell(wsL, f"B{LR['net_debt']}", f"=B{LR['total_debt']}-B{LR['cash']}", F_CALC, NUM)
for c, p in zip(LY, ["B"] + LY[:-1]):
    q = lambda k, col=c: f"{col}{LR[k]}"
    pp = lambda k: f"{p}{LR[k]}"
    f = {
        "revenue": f"={pp('revenue')}*(1+{q('growth')})", "ebitda": f"={q('revenue')}*({q('margin')}-{P('sa_pct')})", "da": f"={q('revenue')}*{P('da_pct')}",
        "ebit": f"={q('ebitda')}-{q('da')}", "capex": f"={q('revenue')}*{P('capex_pct')}", "dnwc": f"={P('nwc_pct')}*({q('revenue')}-{pp('revenue')})",
        "min_cash": f"={P('mincash')}*{q('revenue')}", "mand": f"=MIN({pp('tlb')},{P('amort')}*$B${LI['tlb0']})",
        "int_rcf": (f"=IFERROR(({P('base')}+{P('rcf_sp')})*IF({P('circ')}=1,({pp('revolver')}+{q('revolver')})/2,{pp('revolver')})"
                    f"+{P('undrawn')}*MAX($B${LI['rcf']}-IF({P('circ')}=1,({pp('revolver')}+{q('revolver')})/2,{pp('revolver')}),0),0)"),
        "int_tlb": f"=IFERROR(({P('base')}+{P('tlb_sp')})*IF({P('circ')}=1,({pp('tlb')}+{q('tlb')})/2,{pp('tlb')}),0)", "int_notes": f"={P('coupon')}*{pp('notes_b')}",
        "int_mezz": f"={P('mz_cash')}*{pp('mezz')}", "pik": f"={P('mz_pik')}*{pp('mezz')}",
        "int_inc": f"=IFERROR(MAX({P('base')}-0.005,0)*IF({P('circ')}=1,({pp('cash')}+{q('cash')})/2,{pp('cash')}),0)",
        "cash_int": f"={q('int_rcf')}+{q('int_tlb')}+{q('int_notes')}+{q('int_mezz')}-{q('int_inc')}",
        "gross_int": f"={q('int_rcf')}+{q('int_tlb')}+{q('int_notes')}+{q('int_mezz')}+{q('pik')}+$B${LI['finfees']}/{P('feeyrs')}-{q('int_inc')}",
        "deduct": f"=MIN({q('gross_int')},MAX({P('cap')}*{q('ebitda')},0))", "tax": f"={P('tax')}*MAX({q('ebit')}-{q('deduct')},0)",
        "cfads": f"={q('ebitda')}-{q('tax')}-{q('capex')}-{q('dnwc')}",
        "avail": f"={pp('cash')}+{q('cfads')}-{q('cash_int')}-{q('mand')}-{q('min_cash')}",
        "draw": f"=MIN(MAX(-{q('avail')},0),MAX($B${LI['rcf']}-{pp('revolver')},0))", "repay": f"=MIN({pp('revolver')},MAX({q('avail')},0))",
        "sweep": f"={P('sweep')}*MIN({pp('tlb')}-{q('mand')},MAX({q('avail')}-{q('repay')},0))", "fcf": f"={q('cfads')}-{q('cash_int')}",
        "cash": f"={pp('cash')}+{q('cfads')}-{q('cash_int')}-{q('mand')}+{q('draw')}-{q('repay')}-{q('sweep')}",
        "revolver": f"={pp('revolver')}+{q('draw')}-{q('repay')}", "tlb": f"={pp('tlb')}-{q('mand')}-{q('sweep')}", "notes_b": f"={pp('notes_b')}",
        "mezz": f"={pp('mezz')}+{q('pik')}", "total_debt": f"={q('revolver')}+{q('tlb')}+{q('notes_b')}+{q('mezz')}",
        "net_debt": f"={q('total_debt')}-{q('cash')}", "lev": f"=IF({q('ebitda')}>0,{q('total_debt')}/{q('ebitda')},0)",
        "cover": f"=IF({q('cash_int')}>0,{q('ebitda')}/{q('cash_int')},0)"}
    for k, ff in f.items():
        cell(wsL, q(k), ff, F_CALC, MULT if k in ("lev", "cover") else NUM)
r_ += 1
section(wsL, r_, "Returns at exit"); r_ += 1
rng = lambda k: f"$C${LR[k]}:$I${LR[k]}"
ret_rows = [("ev_x", "Exit EV = exit multiple × EBITDA in the exit year", f"={P('exit')}*INDEX({rng('ebitda')},{P('hold')})", NUM),
            ("nd_x", "Net debt at exit", f"=INDEX({rng('net_debt')},{P('hold')})", NUM),
            ("cost_x", "Exit costs", f"={P('exit_cost')}*B{{ev_x}}", NUM),
            ("eq_pre", "Equity before promote", "=B{ev_x}-B{nd_x}-B{cost_x}", NUM),
            ("prom", "Management promote", f"=MAX(B{{eq_pre}}-$B${LI['equity']},0)*{P('promote')}", NUM),
            ("eq_x", "Sponsor equity at exit", "=B{eq_pre}-B{prom}", NUM),
            ("moic", "Money multiple (MOIC)", f"=B{{eq_x}}/$B${LI['equity']}", MULT),
            ("irr", "IRR", f"=IF(B{{moic}}>0,B{{moic}}^(1/{P('hold')})-1,-1)", PCT)]
for k, *_ in ret_rows:
    LI[k] = r_; r_ += 1
for k, lab, f, fmt in ret_rows:
    cell(wsL, f"A{LI[k]}", lab, bold=k in ("eq_x", "irr"))
    cell(wsL, f"B{LI[k]}", f.format(**{kk: LI[kk] for kk in LI}), F_CALC, fmt, bold=k in ("eq_x", "irr"), fill=FILL_KEY if k == "irr" else None)
widths(wsL, 50, 12, 10)

# ═════════════════════════════ Merger ═════════════════════════════
wsM = wb.create_sheet("Merger", 4)
title(wsM, "Shell acquires BP · deal terms, purchase accounting and EPS", "USD millions. Terms and accounting are live formulas. "
      "Standalone and pro forma net income come from the Python engine (the full pro forma forecast is in model/ma.py) and are marked as such.")
d0 = DEAL["inputs"]
bB = DEAL["B"]["out"]["balance_sheet"][2026]
MI = {}
m_inputs = [("Market data", [("px_s", "Shell share price (GBp)", DEAL["A"]["market"]["price_gbp_pence"], NUM1), ("sh_s", "Shell shares (m)", DEAL["shares"]["shell"], NUM1),
                             ("px_b", "BP share price (GBp)", DEAL["B"]["market"]["price_gbp_pence"], NUM1), ("sh_b", "BP shares (m)", DEAL["B"]["market"]["shares_m"], NUM1),
                             ("fx", "GBP/USD", DEAL["value"]["fx"], NUM4)]),
            ("Terms", [("prem", "Premium to BP share price", d0.premium, PCT), ("stock", "Paid in Shell shares", d0.pct_stock, PCT),
                       ("advp", "Advisory fees, % target EV", d0.advisory_fee_pct, PCT2), ("finp", "Financing fees, % acquisition debt", d0.financing_fee_pct, PCT2),
                       ("cashbs", "Cash consideration funded from cash (USD m)", d0.cash_from_balance_sheet, NUM), ("div", "BP divested, % of EBITDA", d0.divest_pct_ebitda, PCT),
                       ("divm", "Divestment EV/EBITDA", d0.divest_multiple, MULT), ("wu", "Fair-value uplift on BP fixed assets", d0.ppa_writeup_pct, PCT),
                       ("dtl", "Deferred tax rate on uplift", d0.dtl_rate, PCT), ("life", "Uplift depreciation life (years)", d0.writeup_life_years, "0")]),
            ("BP at close, 31 Dec 2026 (Python standalone forecast)", [("b_eq", "Equity attributable to shareholders", bB["equity_parent"], NUM), ("b_gw", "Goodwill", bB["goodwill"], NUM),
                       ("b_fa", "Fixed assets", bB["fixed_assets"], NUM), ("b_borr", "Borrowings + revolver", bB["borrowings"] + bB["revolver"], NUM),
                       ("b_lease", "Lease liabilities", bB["leases"], NUM), ("b_cash", "Cash", bB["cash"], NUM), ("b_nci", "Minorities (incl. hybrids)", bB["minorities"], NUM),
                       ("b_ebitda", "EBITDA 2026", float(DEAL["B"]["out"]["income_statement"].loc["ebitda", 2026]), NUM)])]
r_ = 4
for sec_name, rows_ in m_inputs:
    section(wsM, r_, sec_name); r_ += 1
    for k, lab, x, fmt in rows_:
        cell(wsM, f"A{r_}", lab); cell(wsM, f"B{r_}", float(x), F_IN, fmt, fill=FILL_KEY if k in ("prem", "stock") else None)
        MI[k] = r_; r_ += 1
M = lambda k: f"$B${MI[k]}"
r_ += 1
section(wsM, r_, "Consideration, funding and ownership"); r_ += 1
mcalc = [("mcap_b", "BP market capitalisation", f"={M('sh_b')}*{M('px_b')}/100*{M('fx')}", NUM),
         ("price", "Equity purchase price", "=B{mcap_b}*(1+" + M('prem') + ")", NUM),
         ("offer", "Offer per BP share (GBp)", f"={M('px_b')}*(1+{M('prem')})", NUM1),
         ("stock_v", "Paid in new Shell shares", "=B{price}*" + M('stock'), NUM),
         ("cash_v", "Paid in cash", "=B{price}-B{stock_v}", NUM),
         ("new_sh", "New Shell shares issued (m)", "=B{stock_v}/(" + M('px_s') + "/100*" + M('fx') + ")", NUM1),
         ("pf_sh", "Pro forma shares (m)", "=" + M('sh_s') + "+B{new_sh}", NUM1),
         ("own_bp", "BP holders' ownership", "=B{new_sh}/B{pf_sh}", PCT),
         ("nd_b", "BP net debt incl. leases", f"={M('b_borr')}+{M('b_lease')}-{M('b_cash')}", NUM),
         ("ev_t", "Target enterprise value (price + net debt + minorities)", "=B{price}+B{nd_b}+" + M('b_nci'), NUM),
         ("adv", "Advisory fees", "=" + M('advp') + "*B{ev_t}", NUM),
         ("proc", "Divestment proceeds", f"={M('divm')}*{M('div')}*{M('b_ebitda')}", NUM),
         ("debt", "Acquisition debt (grossed up for its own fee)", "=MAX((B{cash_v}+B{adv}-" + M('cashbs') + "-B{proc})/(1-" + M('finp') + "),0)", NUM),
         ("finfee", "Financing fees", "=" + M('finp') + "*B{debt}", NUM)]
for k, *_ in mcalc:
    MI[k] = r_; r_ += 1
for k, lab, f, fmt in mcalc:
    cell(wsM, f"A{MI[k]}", lab, bold=k in ("price", "debt", "own_bp"))
    cell(wsM, f"B{MI[k]}", f.format(**{kk: MI[kk] for kk in MI}), F_CALC, fmt, bold=k in ("price", "debt", "own_bp"))
r_ += 1
section(wsM, r_, "Purchase price allocation (IFRS 3)"); r_ += 1
ppa_rows = [("wu_v", "Fair-value uplift", f"={M('wu')}*{M('b_fa')}*(1-{M('div')})", NUM), ("dtl_v", "Deferred tax on uplift", "=B{wu_v}*" + M('dtl'), NUM),
            ("fvna", "Fair value of net assets acquired", f"={M('b_eq')}-{M('b_gw')}+B{{wu_v}}-B{{dtl_v}}", NUM),
            ("gw", "Goodwill created", "=B{price}-B{fvna}", NUM), ("stepup", "Annual step-up depreciation", "=B{wu_v}/" + M('life'), NUM)]
for k, *_ in ppa_rows:
    MI[k] = r_; r_ += 1
for k, lab, f, fmt in ppa_rows:
    cell(wsM, f"A{MI[k]}", lab, bold=k == "gw")
    cell(wsM, f"B{MI[k]}", f.format(**{kk: MI[kk] for kk in MI}), F_CALC, fmt, bold=k == "gw")
r_ += 1
section(wsM, r_, "EPS accretion (net income from the Python engine; per-share maths live)"); r_ += 1
yrs = DEAL["years"]
header(wsM, r_, [""] + [f"{y}E" for y in yrs]); r_ += 1
acc = DEAL["accretion"]
er = {}
for k, lab in [("ni_s", "Shell standalone net income to shareholders (Python)"), ("ni_pf", "Pro forma net income to shareholders (Python)"),
               ("eps_s", "Shell standalone EPS (USD)"), ("eps_pf", "Pro forma EPS (USD)"), ("acc", "Accretion / (dilution)")]:
    cell(wsM, f"A{r_}", lab, bold=k == "acc"); er[k] = r_; r_ += 1
for j, y in enumerate(yrs):
    c = COL(2 + j)
    cell(wsM, f"{c}{er['ni_s']}", float(acc.loc["standalone_ni", y]), F_IN, NUM, comment="Python engine output (model/ma.py), same buyback policy as the pro forma.")
    cell(wsM, f"{c}{er['ni_pf']}", float(acc.loc["pro_forma_ni", y]), F_IN, NUM, comment="Python engine output: full pro forma forecast incl. synergies, integration costs, PPA and acquisition debt.")
    cell(wsM, f"{c}{er['eps_s']}", f"={c}{er['ni_s']}/{M('sh_s')}", F_CALC, NUM2)
    cell(wsM, f"{c}{er['eps_pf']}", f"={c}{er['ni_pf']}/$B${MI['pf_sh']}", F_CALC, NUM2)
    cell(wsM, f"{c}{er['acc']}", f"={c}{er['eps_pf']}/{c}{er['eps_s']}-1", F_CALC, PCT, bold=True)
widths(wsM, 56, 13, 7)

# ═════════════════════════════ Portfolio ═════════════════════════════
wsP = wb.create_sheet("Portfolio", 5)
title(wsP, "Portfolio · 13 assets in GBP", "Annualised, from weekly GBP returns since Jan 2021 (Yahoo Finance adjusted closes). Covariance is Ledoit-Wolf shrunk. "
      "Weights come from the Python optimiser (it cannot be expressed as formulas); edit any weight and the statistics recalculate.")
n = len(ASSETS)
r_ = 4
section(wsP, r_, "Inputs"); r_ += 1
PI = {}
for k, lab, x, fmt in [("rf", "Risk-free rate", EST.rf, PCT2), ("erp", "Equity risk premium (CAPM prior)", 0.05, PCT2), ("hw", "Weight on history in expected returns", 0.25, PCT)]:
    cell(wsP, f"A{r_}", lab); cell(wsP, f"B{r_}", float(x), F_IN, fmt, fill=FILL_KEY); PI[k] = r_; r_ += 1
r_ += 1
head_row = r_
strategies = list(PW)
header(wsP, r_, ["Asset", "Hist. mean", "Beta", "Exp. return"] + strategies)
r_ += 1
first = r_
for j, aname in enumerate(ASSETS):
    rr = first + j
    cell(wsP, f"A{rr}", aname)
    cell(wsP, f"B{rr}", float(EST.mu_hist[aname]), F_IN, PCT)
    cell(wsP, f"C{rr}", float(EST.beta[aname]), F_IN, NUM2)
    cell(wsP, f"D{rr}", f"=$B${PI['hw']}*B{rr}+(1-$B${PI['hw']})*($B${PI['rf']}+C{rr}*$B${PI['erp']})", F_CALC, PCT)
    for s_i, s_name in enumerate(strategies):
        cell(wsP, f"{COL(5 + s_i)}{rr}", float(PW[s_name][j]), F_IN, PCT)
lastr = first + n - 1
r_ = lastr + 1
cell(wsP, f"A{r_}", "Sum of weights", F_BOLD)
for s_i in range(len(strategies)):
    c = COL(5 + s_i)
    cell(wsP, f"{c}{r_}", f"=SUM({c}{first}:{c}{lastr})", F_CALC, PCT, bold=True)
sumrow = r_
r_ += 2
cov_top = r_ + 1
section(wsP, r_, "Annual covariance matrix (Ledoit-Wolf)", 2 + n)
header(wsP, cov_top, [""] + ASSETS)
for j, aname in enumerate(ASSETS):
    rr = cov_top + 1 + j
    cell(wsP, f"A{rr}", aname)
    for k_, bname in enumerate(ASSETS):
        cell(wsP, f"{COL(2 + k_)}{rr}", float(EST.cov.loc[aname, bname]), F_IN, NUM4)
cov_rng = f"$B${cov_top + 1}:${COL(1 + n)}${cov_top + n}"
r_ = cov_top + n + 2
# transposed weights (rows) for the quadratic form
section(wsP, r_, "Weights as rows (for w′Σw)", 2 + n); r_ += 1
trow = {}
for s_i, s_name in enumerate(strategies):
    cell(wsP, f"A{r_}", s_name)
    for k_ in range(n):
        cell(wsP, f"{COL(2 + k_)}{r_}", f"={COL(5 + s_i)}{first + k_}", F_CALC, PCT)
    trow[s_name] = r_
    r_ += 1
r_ += 1
section(wsP, r_, "Portfolio statistics", 2 + n); r_ += 1
header(wsP, r_, ["", "Exp. return", "Volatility", "Sharpe", "Beta"]); r_ += 1
PS = {}
for s_i, s_name in enumerate(strategies):
    wc = f"${COL(5 + s_i)}${first}:${COL(5 + s_i)}${lastr}"
    wr = f"$B${trow[s_name]}:${COL(1 + n)}${trow[s_name]}"
    cell(wsP, f"A{r_}", s_name)
    cell(wsP, f"B{r_}", f"=SUMPRODUCT({wc},$D${first}:$D${lastr})", F_CALC, PCT)
    cell(wsP, f"C{r_}", f"=SQRT(SUMPRODUCT({cov_rng}*({wc}*{wr})))", F_CALC, PCT)
    cell(wsP, f"D{r_}", f"=(B{r_}-$B${PI['rf']})/C{r_}", F_CALC, NUM2)
    cell(wsP, f"E{r_}", f"=SUMPRODUCT({wc},$C${first}:$C${lastr})", F_CALC, NUM2)
    PS[s_name] = r_
    r_ += 1
widths(wsP, 26, 12, 3 + n)

# ═════════════════════════════ Checks ═════════════════════════════
wsC = wb.create_sheet("Checks", 6)
title(wsC, "Checks · workbook vs Python engine", "Each line compares a live Excel result with the Python engine. Tolerances are absolute.")
header(wsC, 4, ["Check", "Excel", "Python", "Difference", "Tolerance", "Status"])
dcf_py = r["value_per_share_gbp_pence"]
checks = [
    ("History balances (max abs, FY2021–25)", f"=MAX(ABS(History!B{HCHECK}),ABS(History!C{HCHECK}),ABS(History!D{HCHECK}),ABS(History!E{HCHECK}),ABS(History!F{HCHECK}))", 0.0, 0.5),
    ("Forecast balance sheet balances (max abs)", f"=MAX(Forecast!B{FR['balance_check']}:Forecast!L{FR['balance_check']},-MIN(Forecast!B{FR['balance_check']}:Forecast!L{FR['balance_check']}))", 0.0, 0.01),
    ("Revenue 2035", f"=Forecast!L{FR['revenue']}", float(out["income_statement"].loc["revenue", 2035]), 0.01),
    ("Net income 2026", f"=Forecast!C{FR['net_income']}", float(out["income_statement"].loc["net_income", 2026]), 0.01),
    ("Closing cash 2030", f"=Forecast!G{FR['cash']}", float(out["balance_sheet"].loc["cash", 2030]), 0.01),
    ("Buybacks 2028", f"=Forecast!E{FR['buybacks']}", float(-out["cash_flow"].loc["buybacks", 2028]), 0.01),
    ("Unlevered FCF 2035", f"=Forecast!L{FR['ufcf']}", float(out["kpis"].loc["unlevered_fcf", 2035]), 0.01),
    ("WACC", f"=DCF!B{D['wacc']}", float(r["wacc"]), 1e-9),
    ("Enterprise value", f"=DCF!B{D['ev']}", float(r["ev"]), 0.05),
    ("Equity value", f"=DCF!B{D['equity']}", float(r["equity_value"]), 0.05),
    ("DCF value per share (pence)", f"=DCF!B{D['ps_gbp']}", float(dcf_py), 0.01),
    ("Exit-multiple value per share (pence)", f"=DCF!B{D['ps_m']}", float(r["multiple_method"]["value_per_share_gbp_pence"]), 0.01),
    ("Sensitivity grid centre = point estimate", f"=DCF!{D['sens_centre']}", float(dcf_py), 0.01),
    ("Interest expense 2027", f"=Forecast!D{FR['interest_expense']}", float(out["income_statement"].loc["interest_expense", 2027]), 0.01),
    ("Shareholders' equity 2035", f"=Forecast!L{FR['equity_parent']}", float(out["balance_sheet"].loc["equity_parent", 2035]), 0.01),
    ("LBO sponsor equity", f"=LBO!B{LI['equity']}", float(LBO_R["su"]["equity"]), 0.01),
    ("LBO cash interest, year 1", f"=LBO!C{LR['cash_int']}", float(LBO_R["projections"].loc["cash_interest", 1]), 0.01),
    ("LBO term loan B after year 5", f"=LBO!G{LR['tlb']}", float(LBO_R["projections"].loc["term_loan_b", 5]), 0.01),
    ("LBO IRR", f"=LBO!B{LI['irr']}", float(LBO_R["returns"]["irr"]), 1e-6),
    ("LBO MOIC", f"=LBO!B{LI['moic']}", float(LBO_R["returns"]["moic"]), 1e-6),
    ("Merger: equity purchase price", f"=Merger!B{MI['price']}", float(DEAL["su"]["uses"]["Equity purchase price"]), 0.01),
    ("Merger: acquisition debt", f"=Merger!B{MI['debt']}", float(DEAL["su"]["sources"]["Acquisition debt"]), 0.01),
    ("Merger: BP holders' ownership", f"=Merger!B{MI['own_bp']}", float(DEAL["su"]["ownership_bp"]), 1e-9),
    ("Merger: goodwill created", f"=Merger!B{MI['gw']}", float(DEAL["ppa"]["goodwill_created"]), 0.01),
    (f"Merger: EPS accretion {yrs[2]}", f"=Merger!D{er['acc']}", float(acc.loc["accretion", yrs[2]]), 1e-6),
]
for s_name in strategies:
    st_ = pm.port_stats(PW[s_name], EST)
    checks.append((f"Portfolio volatility · {s_name}", f"=Portfolio!C{PS[s_name]}", st_["vol"], 1e-8))
    checks.append((f"Portfolio expected return · {s_name}", f"=Portfolio!B{PS[s_name]}", st_["ret"], 1e-8))
checks.append(("Portfolio weights sum to 100% (max abs error)", f"=MAX(ABS(Portfolio!E{sumrow}-1),ABS(Portfolio!F{sumrow}-1),ABS(Portfolio!G{sumrow}-1),ABS(Portfolio!H{sumrow}-1),ABS(Portfolio!I{sumrow}-1))", 0.0, 1e-6))
row = 5
for lab, fx_, py, tol in checks:
    cell(wsC, f"A{row}", lab)
    cell(wsC, f"B{row}", fx_, F_LINK, NUM4)
    cell(wsC, f"C{row}", py, F_IN, NUM4)
    cell(wsC, f"D{row}", f"=B{row}-C{row}", F_CALC, NUM4)
    cell(wsC, f"E{row}", tol, F_IN, '0.0E+00')
    cell(wsC, f"F{row}", f'=IFERROR(IF(ABS(D{row})<=E{row},"PASS","FAIL"),"ERROR")', F_CALC, bold=True, align="right")
    row += 1
CHK_LAST = row - 1
cell(wsC, f"A{row + 1}", "Overall", F_BOLD)
cell(wsC, f"F{row + 1}", f'=IF(COUNTIF(F5:F{CHK_LAST},"PASS")=ROWS(F5:F{CHK_LAST}),"ALL PASS",ROWS(F5:F{CHK_LAST})-COUNTIF(F5:F{CHK_LAST},"PASS")&" NOT PASSING")', F_CALC, bold=True, align="right")
CHK_OVERALL = f"F{row + 1}"
widths(wsC, 50, 16, 6)

# ═════════════════════════════ Cover ═════════════════════════════
wsV = wb.create_sheet("Cover", 0)
try:
    from PIL import Image as PImage, ImageDraw
    W_, H_ = 612, 152
    im = PImage.new("RGB", (W_, H_), (255, 255, 255)); dr = ImageDraw.Draw(im)
    sc, ox, oy, wdt = 2.0, 0, 0, 9
    for seg in [((2, 6), (62, 66)), ((62, 6), (2, 66)), ((98, 66), (129, 4)), ((129, 4), (160, 66)), ((196, 6), (227, 71)), ((227, 71), (258, 6)), ((300, 4), (300, 66))]:
        (x1, y1), (x2, y2) = seg
        dr.line([(ox + x1 * sc, oy + y1 * sc), (ox + x2 * sc, oy + y2 * sc)], fill=(17, 17, 20), width=wdt)
    im.save("output/xavi_logo.png")
    logo = XLImage("output/xavi_logo.png"); logo.width, logo.height = 204, 51
    wsV.add_image(logo, "A1")
except Exception as exc:                       # the logo is decorative; never fail the export over it
    cell(wsV, "A1", "XAVI", F_TITLE)
cell(wsV, "A5", "Corporate valuation suite · Shell plc and BP p.l.c.", F_SUB)
cell(wsV, "A6", f"Built {dt.date.today():%d %b %Y} by export_excel.py from the Python engine. Base case. Not investment advice.", F_NOTE)
cell(wsV, "A8", "Model check status", F_BOLD)
cell(wsV, "B8", f"=Checks!{CHK_OVERALL}", F_LINK, bold=True)
cell(wsV, "A10", "Contents", F_SUB)
contents = [("Inputs", "Forecast and valuation assumptions for Shell (edit the blue cells)"), ("Forecast", "Integrated income statement, cash flow and balance sheet, 2026–2035"),
            ("DCF", "WACC, unlevered free cash flow, terminal value, equity bridge, sensitivity grid"), ("LBO", "Carve-out buyout of Shell's Marketing division: sources and uses, debt schedule, returns"),
            ("Merger", "Shell acquires BP: consideration, funding, purchase price allocation, EPS accretion"), ("Portfolio", "13-asset sterling portfolio: expected returns, covariance, portfolio statistics"),
            ("Checks", "Every key result reconciled to the Python engine"), ("History", "Shell FY2021–FY2025 reported figures from the 20-F")]
for j, (s_name, desc) in enumerate(contents):
    c = wsV.cell(row=11 + j, column=1, value=s_name)
    c.hyperlink = f"#'{s_name}'!A1"
    c.font = Font(name=AR, size=10, color="0563C1", underline="single")
    cell(wsV, f"B{11 + j}", desc)
cell(wsV, "A20", "Colour code", F_SUB)
for j, (txt, font, fill) in enumerate([("Blue text: hard-coded input", F_IN, None), ("Black text: formula", F_CALC, None),
                                        ("Green text: link to another sheet", F_LINK, None), ("Yellow fill: key assumption", F_CALC, FILL_KEY)]):
    cell(wsV, f"A{21 + j}", txt, font, fill=fill)
cell(wsV, "A26", "Calculation", F_SUB)
cell(wsV, "A27", "Interest on average balances makes the Forecast and LBO sheets circular. Iterative calculation is saved with the workbook; "
     "if Excel warns about circular references, enable File › Options › Formulas › Enable iterative calculation.", F_NOTE)
wsV.column_dimensions["A"].width = 34
wsV.column_dimensions["B"].width = 90

wb.move_sheet("Cover", -(len(wb.sheetnames) - 1))
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.active = 0
wb.save(OUT)
print(f"wrote {OUT}: sheets {wb.sheetnames}; {len(checks)} reconciliation checks")
