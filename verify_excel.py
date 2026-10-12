"""Read the Excel-recalculated workbook: scan every sheet for formula errors and report the Checks sheet."""
import sys
from openpyxl import load_workbook

PATH = sys.argv[1] if len(sys.argv) > 1 else "output/XAVI_valuation_model.xlsx"
ERRORS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#NUM!", "#N/A", "#NULL!")

vals = load_workbook(PATH, data_only=True)
forms = load_workbook(PATH)
n_formulas, errs, empty = 0, [], []
for ws in forms.worksheets:
    wv = vals[ws.title]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                n_formulas += 1
                v = wv[c.coordinate].value
                if isinstance(v, str) and v in ERRORS:
                    errs.append(f"{ws.title}!{c.coordinate} {v}")
                elif v is None:
                    empty.append(f"{ws.title}!{c.coordinate}")
print(f"formulas: {n_formulas}   errors: {len(errs)}   uncalculated: {len(empty)}")
for e in errs[:30]:
    print("  ERROR", e)
for e in empty[:10]:
    print("  EMPTY", e)

ck = vals["Checks"]
fails = 0
for r in range(5, ck.max_row + 1):
    lab, xl, py, diff, tol, status = (ck.cell(r, c).value for c in range(1, 7))
    if status in ("PASS", "FAIL", "ERROR"):
        fails += status != "PASS"
        print(f"  {status}  {lab:<50} excel={xl!s:>22} python={py!s:>22} diff={diff}")
    elif lab == "Overall":
        print("OVERALL:", status)
sys.exit(1 if (errs or fails or empty) else 0)
